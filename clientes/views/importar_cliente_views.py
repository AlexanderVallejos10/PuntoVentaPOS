import openpyxl

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import redirect, render

from clientes.models import Cliente
from ventas.services.reglas_peru import limpiar_documento


def importar_clientes(request):
    if request.method == 'POST':
        archivo = request.FILES.get('archivo')

        if not archivo:
            messages.error(request, 'Seleccione un archivo Excel.')
            return redirect('clientes:importar_clientes')

        workbook = openpyxl.load_workbook(archivo, data_only=True)
        hoja = workbook.active
        creados = 0
        actualizados = 0

        for fila in hoja.iter_rows(min_row=2, values_only=True):
            nombre, tipo_documento, numero_documento, telefono, email, direccion = (list(fila) + [None] * 6)[:6]
            nombre = str(nombre or '').strip()

            if not nombre:
                continue

            tipo_documento = str(tipo_documento or 'SD').upper().strip() or 'SD'
            numero_documento = limpiar_documento(numero_documento)

            datos = {
                'nombre': nombre,
                'tipo_documento': tipo_documento,
                'numero_documento': numero_documento or '',
                'telefono': str(telefono or '').strip(),
                'email': str(email or '').strip(),
                'direccion': str(direccion or '').strip(),
                'activo': True,
            }

            if numero_documento:
                _, creado = Cliente.objects.update_or_create(
                    tipo_documento=tipo_documento,
                    numero_documento=numero_documento,
                    defaults=datos,
                )
            else:
                Cliente.objects.create(**datos)
                creado = True

            creados += 1 if creado else 0
            actualizados += 0 if creado else 1

        messages.success(request, f'Clientes importados. Nuevos: {creados}. Actualizados: {actualizados}.')
        return redirect('clientes:cliente_list')

    return render(request, 'clientes/importar_clientes.html')


def descargar_ejemplo_clientes(request):
    workbook = openpyxl.Workbook()
    hoja = workbook.active

    hoja.append([
        'nombre',
        'tipo_documento',
        'numero_documento',
        'telefono',
        'email',
        'direccion',
    ])

    hoja.append([
        'JUAN PEREZ',
        'DNI',
        '74413880',
        '999888777',
        'cliente@correo.com',
        'LIMA',
    ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=ejemplo_clientes.xlsx'

    workbook.save(response)
    return response
