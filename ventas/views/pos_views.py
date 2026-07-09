from decimal import Decimal
from io import BytesIO

import openpyxl

from django.contrib import messages
from django.core.mail import EmailMessage, get_connection
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Sum, F, DecimalField, ExpressionWrapper
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, redirect, get_object_or_404

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, letter

from clientes.models import Cliente
from inventario.models import StockBodega, Sede, ConfiguracionEmpresa
from productos.models import Producto, Categoria
from usuarios.decorators import vendedor_required
from usuarios.models import UsuarioSede
from ventas.models import Venta, DetalleVenta, AperturaCaja, MovimientoCaja
from reportlab.lib import colors
from ventas.services.reglas_peru import (
    calcular_totales,
    dinero as money,
    normalizar_tipo_comprobante,
    serie_por_tipo,
    validar_cliente_venta,
)
from ventas.services.pos_service import (
    calcular_subtotal_carrito as subtotal_carrito,
    registrar_detalles_y_stock,
    validar_stock_carrito as validar_stock_pos,
)


def dibujar_cabecera_pdf(pdf):
    empresa = ConfiguracionEmpresa.obtener_configuracion()

    width, height = letter
    y = height - 45

    if empresa and empresa.logo:
        try:
            pdf.drawImage(
                empresa.logo.path,
                40,
                y - 45,
                width=70,
                height=45,
                preserveAspectRatio=True,
                mask='auto'
            )
        except Exception:
            pass

    pdf.setFont('Helvetica-Bold', 16)
    pdf.drawString(
        125,
        y,
        empresa.nombre_empresa if empresa else 'PuntoVentaPOS'
    )

    y -= 18
    pdf.setFont('Helvetica', 9)

    if empresa:
        if empresa.ruc:
            pdf.drawString(125, y, f'RUC: {empresa.ruc}')
            y -= 13

        if empresa.direccion:
            pdf.drawString(125, y, f'Dirección: {empresa.direccion}')
            y -= 13

        if empresa.telefono:
            pdf.drawString(125, y, f'Teléfono: {empresa.telefono}')
            y -= 13

        if empresa.email:
            pdf.drawString(125, y, f'Correo: {empresa.email}')
            y -= 13

    y -= 25

    return y


@vendedor_required
def obtener_carrito(request):
    return request.session.get('carrito', {})


@vendedor_required
def guardar_carrito(request, carrito):
    request.session['carrito'] = carrito
    request.session.modified = True


def usuario_es_admin_o_super(usuario):
    return usuario.is_superuser or usuario.groups.filter(name='Administrador').exists()


# Alias usado por filtrar_historial_ventas y vistas de reportes
def es_admin_o_super(usuario):
    return usuario_es_admin_o_super(usuario)


def obtener_caja_abierta_usuario(usuario):
    return AperturaCaja.objects.filter(
        usuario=usuario,
        estado='ABIERTA'
    ).select_related(
        'caja',
        'caja__sede'
    ).first()


# Selección manual de sede para administradores
def obtener_sede_pos(request):
    usuario = request.user

    if usuario_es_admin_o_super(usuario):
        sede_id = request.session.get('sede_pos_id')
        if sede_id:
            return Sede.objects.filter(id=sede_id, activo=True).first()
        return None

    usuario_sede = UsuarioSede.objects.filter(
        usuario=usuario
    ).select_related('sede').first()

    return usuario_sede.sede if usuario_sede else None


@vendedor_required
def seleccionar_sede_pos(request):
    if not usuario_es_admin_o_super(request.user):
        return redirect('ventas:pos_venta')

    sedes = Sede.objects.filter(activo=True).order_by('nombre')

    if request.method == 'POST':
        sede_id = request.POST.get('sede_id')

        if not sede_id:
            messages.error(request, 'Seleccione una sede.')
            return redirect('ventas:seleccionar_sede_pos')

        request.session['sede_pos_id'] = sede_id
        request.session.modified = True

        return redirect('ventas:pos_venta')

    return render(
        request,
        'ventas/seleccionar_sede_pos.html',
        {'sedes': sedes}
    )


def validar_stock_carrito(carrito, sede):
    return validar_stock_pos(carrito, sede)


def calcular_subtotal_carrito(carrito):
    return subtotal_carrito(carrito)


def calcular_totales_venta(request, subtotal):
    return calcular_totales(
        subtotal=subtotal,
        descuento_porcentaje=request.POST.get('descuento_porcentaje') or 0,
        costo_envio=request.POST.get('costo_envio') or 0,
        cortesia=request.POST.get('cortesia') == 'on',
        envio_domicilio=request.POST.get('envio_domicilio') == 'on',
    )


@vendedor_required
def pos_venta(request):
    apertura_caja = obtener_caja_abierta_usuario(request.user)

    if not apertura_caja:
        messages.error(request, 'Debes abrir caja antes de vender.')
        return redirect('ventas:abrir_caja')

    sede_actual = apertura_caja.caja.sede

    monto_actual_caja = apertura_caja.calcular_monto_esperado()

    movimientos_caja = apertura_caja.movimientos.all().order_by(
        '-created'
    )[:10]

    carrito = obtener_carrito(request)

    productos_carrito = []
    total = Decimal('0.00')

    for producto_id, item in carrito.items():
        producto = get_object_or_404(Producto, id=producto_id)
        cantidad = int(item['cantidad'])

        stock_bodega = StockBodega.objects.filter(
            producto=producto,
            sede=sede_actual,
            activo=True
        ).first()

        producto.stock_pos = stock_bodega.stock if stock_bodega else 0
        subtotal = money(Decimal(str(cantidad)) * producto.precio_venta)

        productos_carrito.append({
            'producto': producto,
            'cantidad': cantidad,
            'subtotal': subtotal,
        })

        total += subtotal

    categoria_id = request.GET.get('categoria')

    stocks = StockBodega.objects.select_related(
        'producto',
        'producto__categoria',
        'sede'
    ).filter(
        sede=sede_actual,
        activo=True,
        producto__activo=True,
        stock__gt=0
    )

    if categoria_id:
        stocks = stocks.filter(producto__categoria_id=categoria_id)

    productos = []

    for stock in stocks.order_by('producto__nombre'):
        producto = stock.producto
        producto.stock_pos = stock.stock
        productos.append(producto)

    categorias = Categoria.objects.filter(activo=True).order_by('nombre')

    ventas_espera = Venta.objects.filter(
        estado='ESPERA'
    ).prefetch_related(
        'detalles__producto'
    ).order_by('id')

    clientes = Cliente.objects.filter(activo=True).order_by('nombre')

    return render(
        request,
        'ventas/pos_venta.html',
        {
            'productos_carrito': productos_carrito,
            'total': money(total),
            'productos': productos,
            'categorias': categorias,
            'ventas_espera': ventas_espera,
            'sede_actual': sede_actual,
            'es_admin_o_super': usuario_es_admin_o_super(request.user),
            'clientes': clientes,
            'apertura_caja': apertura_caja,
            'monto_actual_caja': monto_actual_caja,
            'movimientos_caja': movimientos_caja,
        }
    )


@vendedor_required
def agregar_producto_scanner(request):
    apertura_caja = obtener_caja_abierta_usuario(request.user)

    if not apertura_caja:
        messages.error(request, 'Debes abrir caja antes de vender.')
        return redirect('ventas:abrir_caja')

    sede_actual = apertura_caja.caja.sede

    if request.method == 'POST':
        producto_id = request.POST.get('producto_id')
        codigo = request.POST.get('codigo', '').strip()

        if producto_id:
            producto = Producto.objects.filter(id=producto_id, activo=True).first()
        else:
            if not codigo:
                messages.error(request, 'Debes ingresar o escanear un código.')
                return redirect('ventas:pos_venta')

            producto = Producto.objects.filter(codigo=codigo, activo=True).first()

        if not producto:
            if producto_id:
                messages.error(request, 'El producto seleccionado no existe o está inactivo.')
            else:
                messages.error(request, f'No existe producto con código: {codigo}')
            return redirect('ventas:pos_venta')

        stock_bodega = StockBodega.objects.filter(
            producto=producto,
            sede=sede_actual,
            activo=True
        ).first()

        if not stock_bodega or stock_bodega.stock <= 0:
            messages.error(request, f'{producto.nombre} no tiene stock disponible en {sede_actual.nombre}.')
            return redirect('ventas:pos_venta')

        carrito = obtener_carrito(request)
        producto_id = str(producto.id)

        if producto_id in carrito:
            nueva_cantidad = int(carrito[producto_id]['cantidad']) + 1

            if nueva_cantidad > stock_bodega.stock:
                messages.error(
                    request,
                    f'No hay más stock disponible para {producto.nombre}. '
                    f'Stock en {sede_actual.nombre}: {stock_bodega.stock}.'
                )
                return redirect('ventas:pos_venta')

            carrito[producto_id]['cantidad'] = nueva_cantidad
        else:
            carrito[producto_id] = {'cantidad': 1}

        guardar_carrito(request, carrito)

    return redirect('ventas:pos_venta')


@vendedor_required
@transaction.atomic
def guardar_venta_espera(request):
    apertura_caja = obtener_caja_abierta_usuario(request.user)

    if not apertura_caja:
        messages.error(request, 'Debes abrir caja antes de vender.')
        return redirect('ventas:abrir_caja')

    sede_actual = apertura_caja.caja.sede
    carrito = obtener_carrito(request)

    if not carrito:
        messages.error(request, 'No hay productos para guardar en espera.')
        return redirect('ventas:pos_venta')

    stock_ok, mensaje = validar_stock_carrito(carrito, sede_actual)

    if not stock_ok:
        messages.error(request, mensaje)
        return redirect('ventas:pos_venta')

    subtotal_venta = calcular_subtotal_carrito(carrito)

    venta = Venta.objects.create(
        sede=sede_actual,
        subtotal=subtotal_venta,
        descuento=Decimal('0.00'),
        impuesto=Decimal('0.00'),
        total=subtotal_venta,
        estado='ESPERA',
        metodo_pago='EFECTIVO',
        monto_recibido=Decimal('0.00'),
        cambio=Decimal('0.00')
    )

    for producto_id, item in carrito.items():
        producto = get_object_or_404(Producto, id=producto_id)
        cantidad = int(item['cantidad'])
        subtotal = money(producto.precio_venta * cantidad)

        DetalleVenta.objects.create(
            venta=venta,
            producto=producto,
            cantidad=cantidad,
            precio_unitario=producto.precio_venta,
            subtotal=subtotal
        )

    guardar_carrito(request, {})

    messages.success(request, f'Venta en espera #{venta.id} guardada.')
    return redirect('ventas:pos_venta')


@vendedor_required
def cargar_venta_espera(request, venta_id):
    apertura_caja = obtener_caja_abierta_usuario(request.user)

    if not apertura_caja:
        messages.error(request, 'Debes abrir caja antes de vender.')
        return redirect('ventas:abrir_caja')

    sede_actual = apertura_caja.caja.sede

    venta = get_object_or_404(
        Venta.objects.prefetch_related('detalles__producto'),
        id=venta_id,
        estado='ESPERA'
    )

    carrito = {}

    for detalle in venta.detalles.all():
        stock_bodega = StockBodega.objects.filter(
            producto=detalle.producto,
            sede=sede_actual,
            activo=True
        ).first()

        if not stock_bodega:
            messages.error(request, f'{detalle.producto.nombre} no tiene stock en {sede_actual.nombre}.')
            return redirect('ventas:pos_venta')

        if detalle.cantidad > stock_bodega.stock:
            messages.error(request, f'Stock insuficiente para {detalle.producto.nombre}.')
            return redirect('ventas:pos_venta')

        carrito[str(detalle.producto.id)] = {
            'cantidad': int(detalle.cantidad)
        }

    guardar_carrito(request, carrito)

    request.session['venta_espera_id'] = venta.id
    request.session.modified = True

    messages.success(request, f'Venta en espera #{venta.id} cargada correctamente.')
    return redirect('ventas:pos_venta')


@vendedor_required
def eliminar_venta_espera(request, venta_id):
    venta = get_object_or_404(Venta, id=venta_id, estado='ESPERA')

    venta.delete()

    if request.session.get('venta_espera_id') == venta_id:
        request.session.pop('venta_espera_id', None)
        guardar_carrito(request, {})

    messages.success(request, 'Venta en espera eliminada.')
    return redirect('ventas:pos_venta')


@vendedor_required
@transaction.atomic
def confirmar_venta(request):
    apertura_caja = obtener_caja_abierta_usuario(request.user)

    if not apertura_caja:
        messages.error(request, 'No tienes una caja abierta.')
        return redirect('ventas:abrir_caja')

    sede_actual = apertura_caja.caja.sede
    carrito = obtener_carrito(request)

    if not carrito:
        messages.error(request, 'No hay productos en el carrito.')
        return redirect('ventas:pos_venta')

    stock_ok, mensaje = validar_stock_carrito(carrito, sede_actual)

    if not stock_ok:
        messages.error(request, mensaje)
        return redirect('ventas:pos_venta')

    subtotal_venta = calcular_subtotal_carrito(carrito)
    totales = calcular_totales_venta(request, subtotal_venta)

    metodo_pago = request.POST.get('metodo_pago', 'EFECTIVO')
    monto_recibido = money(request.POST.get('monto_recibido') or 0)
    cliente_id = request.POST.get('cliente_id') or None
    cliente = None

    if cliente_id:
        cliente = Cliente.objects.filter(id=cliente_id, activo=True).first()

        if not cliente:
            messages.error(request, 'El cliente seleccionado no existe o está inactivo.')
            return redirect('ventas:pos_venta')

    tipo_comprobante = normalizar_tipo_comprobante(
        request.POST.get('tipo_comprobante')
    )

    cliente_ok, mensaje_cliente = validar_cliente_venta(
        totales['total'],
        cliente,
        tipo_comprobante,
    )

    if not cliente_ok:
        messages.error(request, mensaje_cliente)
        return redirect('ventas:pos_venta')

    if totales['cortesia']:
        monto_recibido = Decimal('0.00')
        cambio = Decimal('0.00')
    elif metodo_pago == 'EFECTIVO':
        if monto_recibido < totales['total']:
            messages.error(request, 'El monto recibido no puede ser menor al total.')
            return redirect('ventas:pos_venta')

        cambio = money(monto_recibido - totales['total'])
    else:
        monto_recibido = totales['total']
        cambio = Decimal('0.00')

    venta_espera_id = request.session.get('venta_espera_id')

    if venta_espera_id:
        venta = get_object_or_404(Venta, id=venta_espera_id, estado='ESPERA')
        venta.detalles.all().delete()
    else:
        venta = Venta.objects.create(
            sede=sede_actual,
            total=Decimal('0.00'),
            estado='PAGADA'
        )

    try:
        registrar_detalles_y_stock(venta, carrito, sede_actual)
    except ValueError as error:
        transaction.set_rollback(True)
        messages.error(request, str(error))
        return redirect('ventas:pos_venta')

    venta.cliente            = cliente
    venta.sede               = sede_actual
    venta.tipo_comprobante   = tipo_comprobante
    venta.serie_comprobante  = venta.serie_comprobante or serie_por_tipo(tipo_comprobante)
    venta.numero_comprobante = venta.numero_comprobante or f'{venta.id:08d}'
    venta.subtotal           = subtotal_venta
    venta.descuento          = totales['descuento']
    venta.impuesto           = totales['impuesto']
    venta.total              = totales['total']
    venta.metodo_pago        = metodo_pago
    venta.monto_recibido     = monto_recibido
    venta.cambio             = cambio
    venta.cortesia           = totales['cortesia']
    venta.envio_domicilio    = totales['envio_domicilio']
    venta.origen_envio       = request.POST.get('origen_envio') or ''
    venta.destino_envio      = request.POST.get('destino_envio') or ''
    venta.costo_envio        = totales['costo_envio']
    venta.estado             = 'PAGADA'
    venta.save()

    MovimientoCaja.objects.create(
        apertura=apertura_caja,
        venta=venta,
        tipo='INGRESO',
        concepto=f'Venta {venta.comprobante_codigo}',
        monto=venta.total
    )

    guardar_carrito(request, {})

    request.session.pop('venta_espera_id', None)
    request.session.modified = True

    messages.success(request, f'Venta {venta.comprobante_codigo} registrada correctamente.')

    return redirect('ventas:resumen_venta', venta_id=venta.id)


@vendedor_required
def quitar_producto_carrito(request, producto_id):
    carrito = obtener_carrito(request)
    producto_id = str(producto_id)

    if producto_id in carrito:
        del carrito[producto_id]
        guardar_carrito(request, carrito)

    return redirect('ventas:pos_venta')


@vendedor_required
def limpiar_carrito(request):
    guardar_carrito(request, {})

    request.session.pop('venta_espera_id', None)
    request.session.modified = True

    messages.success(request, 'Carrito limpiado correctamente.')
    return redirect('ventas:pos_venta')


@vendedor_required
def aumentar_cantidad(request, producto_id):
    apertura_caja = obtener_caja_abierta_usuario(request.user)

    if not apertura_caja:
        messages.error(request, 'Debes abrir caja antes de vender.')
        return redirect('ventas:abrir_caja')

    sede_actual = apertura_caja.caja.sede
    carrito = obtener_carrito(request)
    producto_id_str = str(producto_id)

    producto = get_object_or_404(Producto, id=producto_id)

    stock_bodega = StockBodega.objects.filter(
        producto=producto,
        sede=sede_actual,
        activo=True
    ).first()

    if not stock_bodega:
        messages.error(request, f'{producto.nombre} no tiene stock en esta sede.')
        return redirect('ventas:pos_venta')

    if producto_id_str in carrito:
        nueva_cantidad = int(carrito[producto_id_str]['cantidad']) + 1

        if nueva_cantidad > stock_bodega.stock:
            messages.error(
                request,
                f'No hay más stock disponible para {producto.nombre}. '
                f'Stock en {sede_actual.nombre}: {stock_bodega.stock}.'
            )
            return redirect('ventas:pos_venta')

        carrito[producto_id_str]['cantidad'] = nueva_cantidad

    guardar_carrito(request, carrito)
    return redirect('ventas:pos_venta')


@vendedor_required
def disminuir_cantidad(request, producto_id):
    carrito = obtener_carrito(request)
    producto_id_str = str(producto_id)

    if producto_id_str in carrito:
        carrito[producto_id_str]['cantidad'] = int(
            carrito[producto_id_str]['cantidad']
        ) - 1

        if carrito[producto_id_str]['cantidad'] <= 0:
            del carrito[producto_id_str]

    guardar_carrito(request, carrito)
    return redirect('ventas:pos_venta')


@vendedor_required
def ticket_venta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.select_related('cliente', 'sede').prefetch_related('detalles__producto'),
        id=venta_id
    )

    return render(request, 'ventas/ticket_venta.html', {'venta': venta})


@vendedor_required
def resumen_venta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.prefetch_related('detalles__producto'),
        id=venta_id
    )

    return render(request, 'ventas/resumen_venta.html', {'venta': venta})


@vendedor_required
def imprimir_ticket(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.select_related('cliente', 'sede').prefetch_related('detalles__producto'),
        id=venta_id
    )

    empresa = ConfiguracionEmpresa.obtener_configuracion()

    return render(request, 'ventas/imprimir_ticket.html', {'venta': venta, 'empresa': empresa})


@vendedor_required
def imprimir_carta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.select_related('cliente', 'sede').prefetch_related('detalles__producto'),
        id=venta_id
    )

    empresa = ConfiguracionEmpresa.obtener_configuracion()

    return render(request, 'ventas/imprimir_carta.html', {'venta': venta, 'empresa': empresa})


# PDF de comprobante
def generar_pdf_venta(venta, empresa):
    buffer = BytesIO()

    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    y = height - 60

    pdf.setFont('Helvetica-Bold', 18)
    pdf.drawString(50, y, empresa.nombre_empresa or 'PuntoVentaPOS')

    y -= 20
    pdf.setFont('Helvetica', 10)
    pdf.drawString(50, y, f'RUC: {empresa.ruc or "-"}')

    y -= 15
    pdf.drawString(50, y, f'Dirección: {empresa.direccion or "-"}')

    y -= 15
    pdf.drawString(50, y, f'Teléfono: {empresa.telefono or "-"}')

    y -= 35
    pdf.setFont('Helvetica-Bold', 14)
    pdf.drawString(50, y, f'{venta.comprobante_nombre.upper()} {venta.comprobante_codigo}')

    y -= 25
    pdf.setFont('Helvetica', 10)
    pdf.drawString(50, y, f'Cliente: {venta.cliente.nombre if venta.cliente else "-"}')

    y -= 15
    pdf.drawString(50, y, f'Fecha: {venta.created.strftime("%d/%m/%Y %H:%M")}')

    y -= 30
    pdf.setFont('Helvetica-Bold', 10)
    pdf.drawString(50, y, 'Producto')
    pdf.drawString(300, y, 'Cantidad')
    pdf.drawString(380, y, 'Precio')
    pdf.drawString(460, y, 'Subtotal')

    y -= 15
    pdf.setFont('Helvetica', 10)

    for detalle in venta.detalles.all():
        pdf.drawString(50, y, detalle.producto.nombre[:35])
        pdf.drawString(310, y, str(detalle.cantidad))
        pdf.drawString(380, y, f'S/ {detalle.precio_unitario}')
        pdf.drawString(460, y, f'S/ {detalle.subtotal}')

        y -= 18

        if y < 80:
            pdf.showPage()
            y = height - 60

    y -= 20
    pdf.setFont('Helvetica-Bold', 11)
    pdf.drawString(380, y, 'Subtotal:')
    pdf.drawString(460, y, f'S/ {venta.subtotal}')

    y -= 15
    pdf.drawString(380, y, 'Descuento:')
    pdf.drawString(460, y, f'S/ {venta.descuento}')

    y -= 15
    pdf.drawString(380, y, 'IGV:')
    pdf.drawString(460, y, f'S/ {venta.impuesto}')

    y -= 15
    pdf.drawString(380, y, 'Total:')
    pdf.drawString(460, y, f'S/ {venta.total}')

    y -= 35
    pdf.setFont('Helvetica', 10)
    pdf.drawString(50, y, empresa.pie_ticket or 'Gracias por su compra')

    pdf.save()
    buffer.seek(0)
    return buffer


# Envío de comprobante por correo
@vendedor_required
def enviar_factura_correo(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.select_related('cliente', 'sede').prefetch_related('detalles__producto'),
        id=venta_id
    )

    empresa = ConfiguracionEmpresa.obtener_configuracion()
    correo_cliente = venta.cliente.email if venta.cliente and venta.cliente.email else ''

    if request.method == 'POST':
        correo_destino = request.POST.get('correo_destino', '').strip()

        if not correo_destino:
            messages.error(request, 'Ingrese el correo del cliente.')
            return redirect('ventas:enviar_factura_correo', venta_id=venta.id)

        if not empresa.smtp_email or not empresa.smtp_password:
            messages.error(request, 'Configure el correo de la empresa en Config > Sucursal.')
            return redirect('ventas:enviar_factura_correo', venta_id=venta.id)

        try:
            connection = get_connection(
                host=empresa.smtp_host,
                port=empresa.smtp_port,
                username=empresa.smtp_email,
                password=empresa.smtp_password,
                use_tls=True
            )

            asunto = f'{venta.comprobante_nombre} {venta.comprobante_codigo} - {empresa.nombre_empresa}'

            cuerpo = f"""
Hola {venta.cliente.nombre if venta.cliente else 'cliente'},

Adjuntamos el resumen de su compra realizada en {empresa.nombre_empresa}.

Detalle de venta:
Comprobante: {venta.comprobante_codigo}
Fecha: {venta.created.strftime('%d/%m/%Y %H:%M')}
Total: S/ {venta.total}
Método de pago: {venta.metodo_pago}

Gracias por su compra.

{empresa.pie_ticket or ''}
"""

            email = EmailMessage(
                subject=asunto,
                body=cuerpo,
                from_email=empresa.smtp_email,
                to=[correo_destino],
                connection=connection
            )

            pdf_buffer = generar_pdf_venta(venta, empresa)
            email.attach(
                f'comprobante_{venta.comprobante_codigo}.pdf',
                pdf_buffer.getvalue(),
                'application/pdf'
            )

            email.send()

            messages.success(request, 'Comprobante enviado correctamente al correo del cliente.')
            return redirect('ventas:resumen_venta', venta_id=venta.id)

        except Exception as error:
            messages.error(request, f'No se pudo enviar el correo: {error}')
            return redirect('ventas:enviar_factura_correo', venta_id=venta.id)

    return render(
        request,
        'ventas/enviar_factura_correo.html',
        {
            'venta': venta,
            'empresa': empresa,
            'correo_cliente': correo_cliente
        }
    )


@vendedor_required
def historial_ventas(request):
    buscar = request.GET.get('buscar', '').strip()
    estado = request.GET.get('estado', '').strip()

    ventas = Venta.objects.select_related(
        'cliente',
        'sede'
    ).prefetch_related(
        'detalles__producto'
    ).order_by(
        '-created'
    )

    if not usuario_es_admin_o_super(request.user):
        ventas = ventas.filter(
            movimientocaja__apertura__usuario=request.user
        ).distinct()

    if estado:
        ventas = ventas.filter(estado=estado)

    if buscar:
        ventas = ventas.filter(
            Q(id__icontains=buscar)
            | Q(serie_comprobante__icontains=buscar)
            | Q(numero_comprobante__icontains=buscar)
            | Q(cliente__nombre__icontains=buscar)
            | Q(metodo_pago__icontains=buscar)
            | Q(sede__nombre__icontains=buscar)
        )

    paginator = Paginator(ventas, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    for venta in page_obj:
        movimiento = MovimientoCaja.objects.filter(
            venta=venta
        ).select_related(
            'apertura',
            'apertura__caja',
            'apertura__usuario'
        ).first()

        venta.movimiento_caja = movimiento

    return render(
        request,
        'ventas/historial_ventas.html',
        {
            'ventas': page_obj,
            'page_obj': page_obj,
            'buscar': buscar,
            'estado': estado,
            'es_admin_o_super': usuario_es_admin_o_super(request.user),
        }
    )


# Filtro común para historial y exportaciones
def filtrar_historial_ventas(request):
    buscar = request.GET.get('buscar', '').strip()
    estado = request.GET.get('estado', '').strip()

    ventas = Venta.objects.select_related(
        'cliente',
        'sede'
    ).order_by(
        '-created'
    )

    if not es_admin_o_super(request.user):
        ventas = ventas.filter(
            movimientocaja__apertura__usuario=request.user
        ).distinct()

    if estado:
        ventas = ventas.filter(
            estado=estado
        )

    if buscar:
        ventas = ventas.filter(
            Q(id__icontains=buscar)
            | Q(serie_comprobante__icontains=buscar)
            | Q(numero_comprobante__icontains=buscar)
            | Q(cliente__nombre__icontains=buscar)
            | Q(metodo_pago__icontains=buscar)
            | Q(sede__nombre__icontains=buscar)
        )

    return ventas


@vendedor_required
def exportar_historial_ventas_excel(request):
    ventas = filtrar_historial_ventas(request)

    workbook = openpyxl.Workbook()
    hoja = workbook.active
    hoja.title = 'Historial ventas'

    hoja.append([
        'Item',
        'Comprobante',
        'Fecha',
        'Cliente',
        'Sede',
        'Caja',
        'Vendedor',
        'Método pago',
        'Total',
        'Estado',
    ])

    for index, venta in enumerate(ventas, start=1):
        movimiento = MovimientoCaja.objects.filter(
            venta=venta
        ).select_related(
            'apertura',
            'apertura__caja',
            'apertura__usuario'
        ).first()

        hoja.append([
            index,
            venta.comprobante_codigo,
            venta.created.strftime('%d/%m/%Y %H:%M'),
            venta.cliente.nombre if venta.cliente else 'Cliente general',
            venta.sede.nombre if venta.sede else '',
            movimiento.apertura.caja.nombre if movimiento else '',
            movimiento.apertura.usuario.username if movimiento else '',
            venta.metodo_pago,
            float(venta.total),
            venta.estado,
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename=historial_ventas.xlsx'

    workbook.save(response)

    return response


@vendedor_required
def exportar_historial_ventas_pdf(request):
    ventas = filtrar_historial_ventas(request)
    empresa = ConfiguracionEmpresa.obtener_configuracion()

    response = HttpResponse(
        content_type='application/pdf'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename=historial_ventas.pdf'

    pdf = canvas.Canvas(
        response,
        pagesize=letter
    )

    width, height = letter
    y = height - 45

    # CABECERA EMPRESA
    if empresa.logo:
        try:
            pdf.drawImage(
                empresa.logo.path,
                40,
                y - 45,
                width=70,
                height=45,
                preserveAspectRatio=True,
                mask='auto'
            )
        except Exception:
            pass

    pdf.setFont('Helvetica-Bold', 16)
    pdf.drawString(
        125,
        y,
        empresa.nombre_empresa or 'PuntoVentaPOS'
    )

    y -= 18
    pdf.setFont('Helvetica', 9)

    if empresa.ruc:
        pdf.drawString(125, y, f'RUC: {empresa.ruc}')
        y -= 13

    if empresa.direccion:
        pdf.drawString(125, y, f'Dirección: {empresa.direccion}')
        y -= 13

    if empresa.telefono:
        pdf.drawString(125, y, f'Teléfono: {empresa.telefono}')
        y -= 13

    if empresa.email:
        pdf.drawString(125, y, f'Correo: {empresa.email}')
        y -= 13

    y -= 22

    pdf.setFont('Helvetica-Bold', 14)
    pdf.drawString(40, y, 'Historial de ventas - Recibos')

    y -= 28

    pdf.setFont('Helvetica-Bold', 7)
    pdf.drawString(40, y, 'Item')
    pdf.drawString(70, y, 'Comprobante')
    pdf.drawString(120, y, 'Fecha')
    pdf.drawString(205, y, 'Cliente')
    pdf.drawString(330, y, 'Sede')
    pdf.drawString(395, y, 'Caja')
    pdf.drawString(445, y, 'Método')
    pdf.drawString(500, y, 'Total')
    pdf.drawString(550, y, 'Estado')

    y -= 15

    pdf.setFont('Helvetica', 7)

    for index, venta in enumerate(ventas, start=1):

        if y < 50:
            pdf.showPage()
            y = height - 45

            pdf.setFont('Helvetica-Bold', 14)
            pdf.drawString(40, y, 'Historial de ventas - Recibos')
            y -= 28

            pdf.setFont('Helvetica-Bold', 7)
            pdf.drawString(40, y, 'Item')
            pdf.drawString(70, y, 'Comprobante')
            pdf.drawString(120, y, 'Fecha')
            pdf.drawString(205, y, 'Cliente')
            pdf.drawString(330, y, 'Sede')
            pdf.drawString(395, y, 'Caja')
            pdf.drawString(445, y, 'Método')
            pdf.drawString(500, y, 'Total')
            pdf.drawString(550, y, 'Estado')
            y -= 15

            pdf.setFont('Helvetica', 7)

        movimiento = MovimientoCaja.objects.filter(
            venta=venta
        ).select_related(
            'apertura',
            'apertura__caja'
        ).first()

        pdf.drawString(40, y, str(index))
        pdf.drawString(70, y, venta.comprobante_codigo)
        pdf.drawString(120, y, venta.created.strftime('%d/%m/%Y %H:%M'))
        pdf.drawString(205, y, (venta.cliente.nombre if venta.cliente else 'Cliente general')[:22])
        pdf.drawString(330, y, (venta.sede.nombre if venta.sede else '-')[:10])
        pdf.drawString(395, y, (movimiento.apertura.caja.nombre if movimiento else '-')[:10])
        pdf.drawString(445, y, venta.metodo_pago[:10])
        pdf.drawString(500, y, f'S/ {venta.total}')
        pdf.drawString(550, y, venta.estado)

        y -= 14

    pdf.save()

    return response


def restaurar_stock_venta(venta):
    for detalle in venta.detalles.select_related('producto').all():
        stock_bodega = StockBodega.objects.filter(
            producto=detalle.producto,
            sede=venta.sede
        ).first()

        if stock_bodega:
            stock_bodega.stock += detalle.cantidad
            stock_bodega.save()


@vendedor_required
@transaction.atomic
def devolver_venta_ajax(request, venta_id):
    if request.method != 'POST':
        return JsonResponse({
            'ok': False,
            'mensaje': 'Método no permitido.'
        })

    venta = get_object_or_404(
        Venta.objects.prefetch_related('detalles__producto'),
        id=venta_id
    )

    if venta.estado in ['ANULADA', 'DEVUELTA']:
        return JsonResponse({
            'ok': False,
            'mensaje': 'La venta ya fue anulada o devuelta.'
        })

    movimiento_original = MovimientoCaja.objects.filter(
        venta=venta,
        tipo='INGRESO'
    ).select_related('apertura').first()

    if not movimiento_original:
        return JsonResponse({
            'ok': False,
            'mensaje': 'No se encontró el movimiento de caja de esta venta.'
        })

    restaurar_stock_venta(venta)

    MovimientoCaja.objects.create(
        apertura=movimiento_original.apertura,
        venta=venta,
        tipo='SALIDA',
        concepto=f'Devolución venta {venta.comprobante_codigo}',
        monto=venta.total
    )

    venta.estado = 'DEVUELTA'
    venta.save()

    return JsonResponse({
        'ok': True,
        'mensaje': 'Venta devuelta correctamente. Stock restaurado y caja actualizada.'
    })


@vendedor_required
@transaction.atomic
def eliminar_venta_ajax(request, venta_id):
    if request.method != 'POST':
        return JsonResponse({
            'ok': False,
            'mensaje': 'Método no permitido.'
        })

    venta = get_object_or_404(
        Venta.objects.prefetch_related('detalles__producto'),
        id=venta_id
    )

    if venta.estado in ['ANULADA', 'DEVUELTA']:
        return JsonResponse({
            'ok': False,
            'mensaje': 'La venta ya fue anulada o devuelta.'
        })

    movimiento_original = MovimientoCaja.objects.filter(
        venta=venta,
        tipo='INGRESO'
    ).select_related('apertura').first()

    restaurar_stock_venta(venta)

    if movimiento_original:
        MovimientoCaja.objects.create(
            apertura=movimiento_original.apertura,
            venta=venta,
            tipo='SALIDA',
            concepto=f'Anulación venta {venta.comprobante_codigo}',
            monto=venta.total
        )

    venta.estado = 'ANULADA'
    venta.save()

    return JsonResponse({
        'ok': True,
        'mensaje': 'Venta anulada correctamente. Stock restaurado y caja ajustada.'
    })


@vendedor_required
@transaction.atomic
def editar_venta(request, venta_id):

    venta = get_object_or_404(
        Venta.objects.prefetch_related('detalles__producto'),
        id=venta_id
    )

    if venta.estado in ['ANULADA', 'DEVUELTA']:
        messages.error(
            request,
            'No se puede editar una venta anulada o devuelta.'
        )
        return redirect('ventas:historial_ventas')

    movimiento_original = MovimientoCaja.objects.filter(
        venta=venta,
        tipo='INGRESO'
    ).select_related('apertura').first()

    if request.method == 'POST':
        total_anterior = venta.total
        producto_ids = request.POST.getlist('producto_id[]')
        cantidades = request.POST.getlist('cantidad[]')
        precios = request.POST.getlist('precio_unitario[]')

        nuevos_detalles = []
        subtotal_nuevo = Decimal('0.00')
        cantidades_actuales = {}

        for detalle in venta.detalles.all():
            producto_id = detalle.producto_id
            cantidades_actuales[producto_id] = cantidades_actuales.get(producto_id, 0) + int(detalle.cantidad)

        for producto_id, cantidad_raw, precio_raw in zip(producto_ids, cantidades, precios):
            if not producto_id:
                continue

            producto = get_object_or_404(Producto, id=producto_id)

            try:
                cantidad = int(cantidad_raw or 0)
                precio = money(precio_raw or 0)
            except (TypeError, ValueError):
                messages.error(request, f'Revise la cantidad o precio de {producto.nombre}.')
                return redirect('ventas:editar_venta', venta_id=venta.id)

            if cantidad <= 0:
                messages.error(request, f'La cantidad de {producto.nombre} debe ser mayor a cero.')
                return redirect('ventas:editar_venta', venta_id=venta.id)

            stock = StockBodega.objects.filter(
                producto=producto,
                sede=venta.sede,
                activo=True,
            ).first()

            disponible = (stock.stock if stock else 0) + cantidades_actuales.get(producto.id, 0)

            if cantidad > disponible:
                messages.error(request, f'Stock insuficiente para {producto.nombre}. Disponible: {disponible}.')
                return redirect('ventas:editar_venta', venta_id=venta.id)

            subtotal = money(precio * cantidad)
            subtotal_nuevo += subtotal
            nuevos_detalles.append((producto, cantidad, precio, subtotal))

        solicitadas = {}
        productos_por_id = {}

        for producto, cantidad, _, _ in nuevos_detalles:
            solicitadas[producto.id] = solicitadas.get(producto.id, 0) + cantidad
            productos_por_id[producto.id] = producto

        for producto_id, cantidad_total in solicitadas.items():
            producto = productos_por_id[producto_id]
            stock = StockBodega.objects.filter(
                producto=producto,
                sede=venta.sede,
                activo=True,
            ).first()
            disponible = (stock.stock if stock else 0) + cantidades_actuales.get(producto_id, 0)

            if cantidad_total > disponible:
                messages.error(request, f'Stock insuficiente para {producto.nombre}. Disponible: {disponible}.')
                return redirect('ventas:editar_venta', venta_id=venta.id)

        if not nuevos_detalles:
            messages.error(request, 'La venta debe tener al menos un producto.')
            return redirect('ventas:editar_venta', venta_id=venta.id)

        descuento = money(request.POST.get('descuento') or 0)
        impuesto = money(request.POST.get('impuesto') or 0)
        costo_envio = money(request.POST.get('costo_envio') or 0)
        metodo_pago = request.POST.get('metodo_pago') or venta.metodo_pago
        total_nuevo = money(subtotal_nuevo - descuento + impuesto + costo_envio)

        cliente_ok, mensaje_cliente = validar_cliente_venta(
            total_nuevo,
            venta.cliente,
            venta.tipo_comprobante,
        )

        if not cliente_ok:
            messages.error(request, mensaje_cliente)
            return redirect('ventas:editar_venta', venta_id=venta.id)

        monto_recibido = money(request.POST.get('monto_recibido') or total_nuevo)

        if metodo_pago == 'EFECTIVO' and monto_recibido < total_nuevo:
            messages.error(request, 'El monto recibido no puede ser menor al total.')
            return redirect('ventas:editar_venta', venta_id=venta.id)

        for detalle in venta.detalles.select_related('producto').all():
            stock = StockBodega.objects.filter(
                producto=detalle.producto,
                sede=venta.sede,
                activo=True,
            ).first()
            if stock:
                stock.stock += detalle.cantidad
                stock.save(update_fields=['stock'])

        venta.detalles.all().delete()

        for producto, cantidad, precio, subtotal in nuevos_detalles:
            DetalleVenta.objects.create(
                venta=venta,
                producto=producto,
                cantidad=cantidad,
                precio_unitario=precio,
                subtotal=subtotal,
            )

            stock = StockBodega.objects.select_for_update().filter(
                producto=producto,
                sede=venta.sede,
                activo=True,
            ).first()
            stock.stock -= cantidad
            stock.save(update_fields=['stock'])

        venta.subtotal = money(subtotal_nuevo)
        venta.descuento = descuento
        venta.impuesto = impuesto
        venta.costo_envio = costo_envio
        venta.total = total_nuevo
        venta.metodo_pago = metodo_pago
        venta.monto_recibido = monto_recibido
        venta.cambio = money(monto_recibido - total_nuevo)
        venta.save()

        diferencia = total_nuevo - total_anterior

        if movimiento_original and diferencia != 0:
            if diferencia > 0:
                tipo = 'INGRESO'
                monto = diferencia
                concepto = f'Aumento por edición venta {venta.comprobante_codigo}'
            else:
                tipo = 'SALIDA'
                monto = abs(diferencia)
                concepto = f'Disminución por edición venta {venta.comprobante_codigo}'

            MovimientoCaja.objects.create(
                apertura=movimiento_original.apertura,
                venta=venta,
                tipo=tipo,
                concepto=concepto,
                monto=monto,
            )

        messages.success(request, 'Venta editada correctamente. Stock y caja actualizados.')
        return redirect('ventas:historial_ventas')

    productos = Producto.objects.all().order_by('nombre')

    return render(
        request,
        'ventas/editar_venta.html',
        {
            'venta': venta,
            'productos': productos
        }
    )


# Detalle de ventas
@vendedor_required
def detalle_ventas(request):
    buscar = request.GET.get('buscar', '').strip()

    detalles = DetalleVenta.objects.select_related(
        'venta',
        'venta__cliente',
        'venta__sede',
        'producto',
        'producto__categoria'
    ).order_by(
        '-venta__created'
    )

    if not es_admin_o_super(request.user):
        detalles = detalles.filter(
            venta__movimientocaja__apertura__usuario=request.user
        ).distinct()

    if buscar:
        detalles = detalles.filter(
            Q(venta__id__icontains=buscar)
            | Q(venta__serie_comprobante__icontains=buscar)
            | Q(venta__numero_comprobante__icontains=buscar)
            | Q(producto__nombre__icontains=buscar)
            | Q(venta__cliente__nombre__icontains=buscar)
            | Q(venta__sede__nombre__icontains=buscar)
        )

    paginator = Paginator(
        detalles,
        10
    )

    page_obj = paginator.get_page(
        request.GET.get('page')
    )

    return render(
        request,
        'ventas/detalle_ventas.html',
        {
            'detalles': page_obj,
            'page_obj': page_obj,
            'buscar': buscar
        }
    )


# Ventas agrupadas por categoría
@vendedor_required
def ventas_por_categoria(request):
    categoria_id = request.GET.get('categoria', '').strip()

    categorias = Categoria.objects.all().order_by(
        'nombre'
    )

    detalles = DetalleVenta.objects.select_related(
        'venta',
        'producto',
        'producto__categoria'
    ).filter(
        venta__estado='PAGADA'
    )

    if not es_admin_o_super(request.user):
        detalles = detalles.filter(
            venta__movimientocaja__apertura__usuario=request.user
        ).distinct()

    if categoria_id:
        detalles = detalles.filter(
            producto__categoria_id=categoria_id
        )

    ventas_categoria = detalles.values(
        'producto__categoria__nombre',
        'producto__nombre'
    ).annotate(
        cantidad_total=Sum('cantidad'),
        total_vendido=Sum('subtotal')
    ).order_by(
        'producto__categoria__nombre',
        '-total_vendido'
    )

    return render(
        request,
        'ventas/ventas_por_categoria.html',
        {
            'categorias': categorias,
            'ventas_categoria': ventas_categoria,
            'categoria_id': categoria_id
        }
    )


@vendedor_required
def exportar_detalle_ventas_excel(request):

    buscar = request.GET.get('buscar', '').strip()

    detalles = DetalleVenta.objects.select_related(
        'venta',
        'venta__cliente',
        'venta__sede',
        'producto',
        'producto__categoria'
    )

    if not es_admin_o_super(request.user):
        detalles = detalles.filter(
            venta__movimientocaja__apertura__usuario=request.user
        ).distinct()

    if buscar:
        detalles = detalles.filter(
            Q(producto__nombre__icontains=buscar)
            | Q(venta__cliente__nombre__icontains=buscar)
            | Q(venta__id__icontains=buscar)
            | Q(venta__serie_comprobante__icontains=buscar)
            | Q(venta__numero_comprobante__icontains=buscar)
        )

    workbook = openpyxl.Workbook()

    hoja = workbook.active
    hoja.title = 'Detalle Ventas'

    hoja.append([
        'Comprobante',
        'Sede',
        'Caja',
        'Fecha',
        'Cliente',
        'Producto',
        'Categoria',
        'Precio',
        'Cantidad',
        'Total',
    ])

    for detalle in detalles:

        movimiento = detalle.venta.movimientocaja_set.first()

        caja = '-'

        if movimiento:
            caja = movimiento.apertura.caja.nombre

        hoja.append([
            detalle.venta.comprobante_codigo,
            detalle.venta.sede.nombre if detalle.venta.sede else '-',
            caja,
            detalle.venta.created.strftime('%d/%m/%Y %H:%M'),
            detalle.venta.cliente.nombre if detalle.venta.cliente else 'Cliente general',
            detalle.producto.nombre,
            detalle.producto.categoria.nombre,
            float(detalle.precio_unitario),
            detalle.cantidad,
            float(detalle.subtotal),
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename=detalle_ventas.xlsx'

    workbook.save(response)

    return response


@vendedor_required
def exportar_detalle_ventas_pdf(request):

    buscar = request.GET.get('buscar', '').strip()

    detalles = DetalleVenta.objects.select_related(
        'venta',
        'venta__cliente',
        'venta__sede',
        'producto',
        'producto__categoria'
    )

    if not es_admin_o_super(request.user):
        detalles = detalles.filter(
            venta__movimientocaja__apertura__usuario=request.user
        ).distinct()

    if buscar:
        detalles = detalles.filter(
            Q(producto__nombre__icontains=buscar)
            | Q(venta__cliente__nombre__icontains=buscar)
            | Q(venta__id__icontains=buscar)
            | Q(venta__serie_comprobante__icontains=buscar)
            | Q(venta__numero_comprobante__icontains=buscar)
        )

    response = HttpResponse(content_type='application/pdf')

    response[
        'Content-Disposition'
    ] = 'attachment; filename=detalle_ventas.pdf'

    pdf = canvas.Canvas(response, pagesize=letter)

    width, height = letter

    y = dibujar_cabecera_pdf(pdf)

    pdf.setFont('Helvetica-Bold', 16)
    pdf.drawString(40, y, 'Detalle de ventas')

    y -= 30

    pdf.setFont('Helvetica-Bold', 8)

    headers = ['Comprobante', 'Producto', 'Cliente', 'Cant', 'Precio', 'Total']
    x_positions = [40, 100, 220, 380, 430, 500]

    for i, header in enumerate(headers):
        pdf.drawString(x_positions[i], y, header)

    y -= 18

    pdf.setFont('Helvetica', 8)

    for detalle in detalles:

        if y < 60:
            pdf.showPage()
            y = dibujar_cabecera_pdf(pdf)

        cliente = 'General'

        if detalle.venta.cliente:
            cliente = detalle.venta.cliente.nombre

        pdf.drawString(40, y, detalle.venta.comprobante_codigo[:12])
        pdf.drawString(100, y, detalle.producto.nombre[:24])
        pdf.drawString(220, y, cliente[:20])
        pdf.drawString(380, y, str(detalle.cantidad))
        pdf.drawString(430, y, f'S/ {detalle.precio_unitario}')
        pdf.drawString(500, y, f'S/ {detalle.subtotal}')

        y -= 16

    pdf.save()

    return response


@vendedor_required
def exportar_ventas_categoria_excel(request):

    categoria_id = request.GET.get('categoria', '')

    detalles = DetalleVenta.objects.select_related(
        'producto',
        'producto__categoria'
    )

    if categoria_id:
        detalles = detalles.filter(producto__categoria_id=categoria_id)

    resumen = detalles.values(
        'producto__categoria__nombre',
        'producto__nombre'
    ).annotate(
        cantidad=Sum('cantidad'),
        total=Sum('subtotal')
    )

    workbook = openpyxl.Workbook()

    hoja = workbook.active
    hoja.title = 'Ventas Categoria'

    hoja.append(['Categoria', 'Producto', 'Cantidad', 'Total Vendido'])

    for item in resumen:
        hoja.append([
            item['producto__categoria__nombre'],
            item['producto__nombre'],
            item['cantidad'],
            float(item['total']),
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response[
        'Content-Disposition'
    ] = 'attachment; filename=ventas_categoria.xlsx'

    workbook.save(response)

    return response


@vendedor_required
def exportar_ventas_categoria_pdf(request):

    categoria_id = request.GET.get('categoria', '')

    detalles = DetalleVenta.objects.select_related(
        'producto',
        'producto__categoria'
    )

    if categoria_id:
        detalles = detalles.filter(producto__categoria_id=categoria_id)

    resumen = detalles.values(
        'producto__categoria__nombre',
        'producto__nombre'
    ).annotate(
        cantidad=Sum('cantidad'),
        total=Sum('subtotal')
    )

    response = HttpResponse(content_type='application/pdf')

    response[
        'Content-Disposition'
    ] = 'attachment; filename=ventas_categoria.pdf'

    pdf = canvas.Canvas(response, pagesize=letter)

    y = dibujar_cabecera_pdf(pdf)

    pdf.setFont('Helvetica-Bold', 16)
    pdf.drawString(40, y, 'Ventas por categoria')

    y -= 30

    pdf.setFont('Helvetica-Bold', 9)
    pdf.drawString(40, y, 'Categoria')
    pdf.drawString(170, y, 'Producto')
    pdf.drawString(380, y, 'Cantidad')
    pdf.drawString(470, y, 'Total')

    y -= 20

    pdf.setFont('Helvetica', 8)

    for item in resumen:

        if y < 60:
            pdf.showPage()
            y = dibujar_cabecera_pdf(pdf)

        pdf.drawString(40, y, item['producto__categoria__nombre'][:18])
        pdf.drawString(170, y, item['producto__nombre'][:30])
        pdf.drawString(380, y, str(item['cantidad']))
        pdf.drawString(470, y, f"S/ {item['total']}")

        y -= 16

    pdf.save()

    return response
